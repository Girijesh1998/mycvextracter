from django.shortcuts import render, redirect

# Create your views here.
# extractor/views.py
import re
import docx
from django.http import HttpResponse
from .forms import CVForm
from .models import CV
import openpyxl

def extract_cv_data(doc):
    text = []
    for para in doc.paragraphs:
        text.append(para.text)
    
    text = '\n'.join(text)
    
    email = re.findall(r'[\w\.-]+@[\w\.-]+', text)
    contact_number = re.findall(r'(\+\d{1,2}\s)?\(?\d{3}\)?[\s.-]?\d{3}[\s.-]?\d{4}', text)
    
    return email[0] if email else None, contact_number[0] if contact_number else None, text

def upload_cv(request):
    if request.method == 'POST':
        form = CVForm(request.POST, request.FILES)
        if form.is_valid():
            cv = form.save(commit=False)
            
            if cv.file.name.endswith('.docx'):
                doc = docx.Document(cv.file)
                email, contact_number, text = extract_cv_data(doc)
                
                cv.email = email
                cv.contact_number = contact_number
                cv.text = text
                cv.save()
                
                # Create Excel file
                workbook = openpyxl.Workbook()
                sheet = workbook.active
                sheet.title = 'CV Data'
                sheet.append(['Email', 'Contact Number', 'Text'])
                sheet.append([email, contact_number, text])
                
                response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
                response['Content-Disposition'] = 'attachment; filename=extracted_cv_data.xlsx'
                workbook.save(response)
                
                return response
    else:
        form = CVForm()
    
    return render(request, 'extract\cv_upload.html', {'form': form})
